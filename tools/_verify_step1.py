import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
import os
files = [
    r'F:\knowledge-base\分析方法验证\验证报告输出\QR-MV2026025-MethodB_step1_字段替换.docx',
    r'F:\knowledge-base\分析方法验证\验证报告输出\QR-MV2026018-MethodB_step1_字段替换.docx',
]
for f in files:
    print(f'=== {f.split(chr(92))[-1]} ===')
    if not os.path.exists(f):
        print('  不存在')
        continue
    size = os.path.getsize(f)
    print(f'  size: {size} bytes')
    p = Presentation(f)
    print(f'  slides: {len(p.slides)}')
    for i, s in enumerate(p.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip()[:80]
                print(f'    slide {i}: {t}')
                break
    print()