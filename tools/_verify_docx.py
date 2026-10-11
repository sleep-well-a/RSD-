import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
import os
files = [
    r'F:\knowledge-base\分析方法验证\验证报告输出\QR-MV2026025-MethodB_step1_字段替换.docx',
    r'F:\knowledge-base\分析方法验证\验证报告输出\QR-MV2026018-MethodB_step1_字段替换.docx',
    r'F:\knowledge-base\分析方法验证\验证报告输出\QR-MV2026029-MethodB_step1_字段替换.docx',
]
for f in files:
    name = f.split(chr(92))[-1]
    print(f'=== {name} ===')
    if not os.path.exists(f):
        print('  不存在')
        continue
    size = os.path.getsize(f)
    print(f'  size: {size} bytes ({size//1024} KB)')
    d = Document(f)
    print(f'  paragraphs: {len(d.paragraphs)}')
    print(f'  tables: {len(d.tables)}')
    # 抽非空段落的文本
    nonempty = [p.text.strip() for p in d.paragraphs if p.text.strip()][:8]
    for j, t in enumerate(nonempty, 1):
        print(f'  p{j}: {t[:80]}')
    print()