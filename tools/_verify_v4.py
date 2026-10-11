import sys, os
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
p = Presentation(r'F:\自制软件\tools\_test_rsd_v54.284_v4.pptx')
print(f'=== v4 verification ===')
print(f'slides: {len(p.slides)}')
print(f'size: {os.path.getsize(r"F:\\自制软件\\tools\\_test_rsd_v54.284_v4.pptx") // 1024} KB')
print()
for i, slide in enumerate(p.slides, 1):
    layout = slide.slide_layout.name
    print(f'Slide {i}: layout="{layout}"')
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip().replace(chr(10), ' | ')[:80]
            if text:
                print(f'  - {text}')
    print()