import sys, os
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
p = Presentation(r'F:\自制软件\tools\_test_rsd_v54.284_t12.pptx')
print(f'=== RSD v54.284 PPT (模板 12) ===')
print(f'幻灯片: {len(p.slides)}')
print(f'大小: {os.path.getsize(r"F:\\自制软件\\tools\\_test_rsd_v54.284_t12.pptx") // 1024} KB')
print()
for i, slide in enumerate(p.slides, 1):
    print(f'Slide {i}:')
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            text = shape.text_frame.text.strip().replace(chr(10), ' | ')[:80]
            if text and not text.startswith('在此'):
                print(f'  - {text}')
    print()